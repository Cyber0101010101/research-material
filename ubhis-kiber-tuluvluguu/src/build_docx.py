"""Markdown төлөвлөгөөг албан бичгийн хэлбэртэй .docx болгон хөрвүүлнэ.

Ажиллуулах: python3 src/build_docx.py  (ubhis-kiber-tuluvluguu хавтаснаас)
Шаардлага: pandoc, python-docx
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "ubhis_kiber_ayulgui_baidal_tuluvluguu.md"
OUT = ROOT / "ubhis_kiber_ayulgui_baidal_tuluvluguu.docx"
FONT = "Times New Roman"


def set_font(style, size, bold=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def add_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def set_col_widths(table, total_cm=16.5):
    """Баганын өргөнийг агуулгын урт, хамгийн урт үгэнд тохируулж хуваарилна."""
    ncols = len(table.columns)
    mins, weights = [], []
    for c in range(ncols):
        texts = [row.cells[c].text for row in table.rows]
        longest_word = max((len(w) for tx in texts for w in tx.split()), default=1)
        mins.append(min(longest_word * 0.2 + 0.35, 4.0))
        weights.append(sum(len(tx) for tx in texts) / len(texts) + 1)
    # Эхлээд хамгийн бага өргөнийг өгч, үлдсэнийг агуулгын уртаар хуваана
    spare = total_cm - sum(mins)
    if spare < 0:
        widths = [m * total_cm / sum(mins) for m in mins]
    else:
        widths = [m + spare * w / sum(weights) for m, w in zip(mins, weights)]
    widths = [Cm(w) for w in widths]
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def main():
    subprocess.run(
        ["pandoc", str(SRC), "-f", "markdown", "-t", "docx", "-o", str(OUT)],
        check=True,
    )
    doc = Document(OUT)

    for sec in doc.sections:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(1.5)
        sec.top_margin = sec.bottom_margin = Cm(2.0)

    by_name = {st.name.lower(): st for st in doc.styles}
    for name in ("normal", "body text", "first paragraph", "compact", "block text"):
        if name in by_name:
            set_font(by_name[name], 12)
    for name, size in (("title", 14), ("heading 1", 13), ("heading 2", 12), ("heading 3", 12)):
        if name in by_name:
            set_font(by_name[name], size, bold=True)
    if "source code" in by_name:
        by_name["source code"].font.size = Pt(9)

    before_title = True
    for p in doc.paragraphs:
        style = p.style.name.lower()
        if style in ("title", "heading 1"):
            before_title = False
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif before_title:
            # Нууцын зэрэглэл, БАТЛАВ хэсэг баруун талд
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif "___" in p.text:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif style in ("body text", "first paragraph", "normal"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        add_borders(table)
        set_col_widths(table)
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.font.name = FONT
                        if i == 0:
                            run.font.bold = True

    doc.save(OUT)
    print(f"Үүсгэлээ: {OUT}")


if __name__ == "__main__":
    main()
